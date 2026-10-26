# -*- coding: utf-8 -*-
"""Convert docs/SRS/srs.md (Markdown) to docs/SRS/srs.docx.

Usage:  python scripts/md_to_docx.py [input.md] [output.docx]
Defaults to docs/SRS/srs.md -> docs/SRS/srs.docx.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

IMG_WIDTH = Cm(15)


def add_runs(paragraph, text):
    """Parse **bold**, *italic* and `code` spans and add them as runs."""
    token_re = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")
    for part in token_re.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 1:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
        else:
            paragraph.add_run(part)


def is_separator_row(cells):
    return all(re.fullmatch(r":?-{2,}:?", c.strip()) for c in cells)


def add_table(doc, rows):
    """rows: list of lists of cell strings; first row is the header."""
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=len(rows), cols=len(header))
    table.style = "Table Grid"
    for r, cells in enumerate(rows):
        for c, cell_text in enumerate(cells):
            # <br> inside cells becomes a line break
            cell = table.cell(r, c)
            cell.text = ""
            lines = cell_text.split("<br>")
            for i, line in enumerate(lines):
                paragraph = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
                if line.strip() == "":
                    continue
                add_runs(paragraph, line.strip())
                if r == 0:
                    for run in paragraph.runs:
                        run.bold = True
    return table


def convert(md_path: Path, out_path: Path):
    doc = Document()

    # Default font (Calibri + Microsoft YaHei for any CJK text)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia",
        "微软雅黑",
    )

    lines = md_path.read_text(encoding="utf-8").splitlines()
    base = md_path.parent
    i = 0
    n_images = n_tables = 0
    while i < len(lines):
        line = lines[i]

        # Image
        m = re.match(r"!\[(.*?)\]\((.*?)\)", line)
        if m:
            img_path = base / m.group(2)
            if img_path.exists():
                doc.add_picture(str(img_path), width=IMG_WIDTH)
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                n_images += 1
            else:
                doc.add_paragraph(f"[missing image: {m.group(2)}]")
            i += 1
            continue

        # Markdown table (consecutive |...| lines)
        if line.strip().startswith("|") and line.strip().endswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                cells = [c for c in lines[i].strip().strip("|").split("|")]
                if not is_separator_row(cells):
                    rows.append(cells)
                i += 1
            if rows:
                add_table(doc, rows)
                n_tables += 1
            continue

        stripped = line.strip()

        # Headings
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=3)
        elif stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=2)
        elif stripped.startswith("# "):
            doc.add_heading(stripped[2:], level=1)
        # Blockquote
        elif stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            add_runs(p, stripped.lstrip("> ").strip())
            for run in p.runs:
                run.italic = True
        # Lists
        elif re.match(r"^- ", stripped):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, stripped[2:])
        elif re.match(r"^\d+\. ", stripped):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\d+\. ", "", stripped))
        # Horizontal rule / blank
        elif stripped in ("", "---", "***"):
            pass
        # Normal paragraph
        else:
            p = doc.add_paragraph()
            add_runs(p, stripped)

        i += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    print(f"OK: {out_path}")
    print(f"  paragraphs={len(doc.paragraphs)}, tables={len(doc.tables)}, images={n_images}, md-tables={n_tables}")


if __name__ == "__main__":
    md = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("docs/SRS/srs.md")
    out = Path(sys.argv[2]) if len(sys.argv) > 2 else md.with_suffix(".docx")
    convert(md, out)
